from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook

#Abre Arquivo Excel
nomeArquivExcel = "DadosAlunos.xlsx"
planilhaDadosAlunos = load_workbook(nomeArquivExcel)

#Selecionando a aba
sheetSelecionada = planilhaDadosAlunos["Nomes"]

for linha in range(2,len(sheetSelecionada["A"]) + 1):

    #Abre Arquivo Word
    arquivoWord = Document("Certificado3.docx")

    #Seleciona o estilo
    estilo = arquivoWord.styles["Normal"]

    #Recuperar nome aluno
    nomeAluno = sheetSelecionada['A%s' % linha].value
    dia = sheetSelecionada['B%s' % linha].value
    mes = sheetSelecionada['C%s' % linha].value
    ano = sheetSelecionada['D%s' % linha].value
    nomeCurso = sheetSelecionada['E%s' % linha].value
    nomeInstrutor = sheetSelecionada['F%s' % linha].value

    for paragrafo in arquivoWord.paragraphs:

        if "@nome" in paragrafo.text:
            paragrafo.text = nomeAluno
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

        paragrafoP1 = "Concluiu com sucesso o curso de "
        paragrafoP2 = ", como carga horária de 20 horas, promovido pela escola de Cursos Online em "
        paragrafoCompleto = paragrafoP1 + nomeCurso + paragrafoP2 + str(dia) + " de " + mes + " de " + str(ano)

        if "escola" in paragrafo.text:
            paragrafo.text = paragrafoCompleto
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

        if "Instrutor" in paragrafo.text:
            paragrafo.text = nomeInstrutor + " - Instrutor"
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

    #Recuperar caminho pasta certificados
    caminhoCertificados = "C:\\Users\\lucas\\Documents\\Curso RPA Python\\pythonProject\\Word\\Certificados\\" + nomeAluno + ".docx"

    #Salvando o arquivo
    arquivoWord.save(caminhoCertificados)

print("Certificados gerados com sucesso")