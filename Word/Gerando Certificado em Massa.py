from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook
import os

#Abre Arquivo Excel
nomeArquivExcel = "Alunos.xlsx"
planilhaDadosAlunos = load_workbook(nomeArquivExcel)

#Selecionando a aba
sheetSelecionada = planilhaDadosAlunos["Nomes"]

for linha in range(2,len(sheetSelecionada["A"]) + 1):

    #Abre Arquivo Word
    arquivoWord = Document("Certificado1.docx")

    #Seleciona o estilo
    estilo = arquivoWord.styles["Normal"]

    #Recuperar nome aluno
    nomeAluno = sheetSelecionada['A%s' % linha].value

    for paragrafo in arquivoWord.paragraphs:

        if "@nome" in paragrafo.text:
            paragrafo.text = nomeAluno
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

    #Recuperar caminho pasta certificados
    caminhoCertificados = "C:\\Users\\lucas\\Documents\\Curso RPA Python\\pythonProject\\Word\\Certificados\\" + nomeAluno + ".docx"

    #Salvando o arquivo
    arquivoWord.save(caminhoCertificados)

print("Certificados gerados com sucesso")